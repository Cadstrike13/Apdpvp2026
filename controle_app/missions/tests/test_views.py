import os
import tempfile
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models.fields.files import FieldFile
from django.test import TestCase, override_settings
from django.urls import reverse

from agents.models import AgentControleur
from entites.models import EntiteControlee
from personnes.models import Personne

from ..models import ControleEntite, EvaluationConformite, MembreGroupeControle, MissionControle, RoleMission, StatutMission

MEDIA_ROOT_TEST = tempfile.mkdtemp()


def creer_controle(**kwargs):
    entite = kwargs.pop("entite_controlee", None) or EntiteControlee.objects.create(nom="ACME SA")
    mission = kwargs.pop("mission", None) or MissionControle.objects.create(date_mission="2026-07-01")
    return ControleEntite.objects.create(mission=mission, entite=entite, **kwargs)


def creer_utilisateur_agent(username, external_id):
    user = User.objects.create_user(username=username, password="pass-123")
    agent = AgentControleur.objects.create(external_id=external_id, nom=username, prenom=username, user=user)
    return user, agent


def declarer_tous_traitements(controle):
    """Coche 'Déclaration effectuée' pour les 10 traitements — équivalent à
    passer par la checkliste (questionnaire_checklist) avant les pages 1 à 5,
    qui ne détaillent que les traitements déclarés."""
    for reponse in controle.reponses.select_related("page1"):
        reponse.page1.declaration_effectuee = True
        reponse.page1.save()


def completer_page(client, controle, page):
    """Soumet la page `page` du questionnaire telle quelle (valeurs par défaut) —
    reproduit le pattern GET (formset+queryset) -> POST utilisé par la vue."""
    r = client.get(reverse("missions:questionnaire_page", args=[controle.pk, page]))
    formset = r.context["formset"]
    data = {f"form-{k}": v for k, v in formset.management_form.initial.items()}
    for i, form in enumerate(formset.forms):
        for name in form.fields:
            value = form.initial.get(name, form.fields[name].initial)
            if value is None or isinstance(value, FieldFile):
                value = ""
            data[f"form-{i}-{name}"] = value
        data[f"form-{i}-id"] = form.instance.pk
    return client.post(reverse("missions:questionnaire_page", args=[controle.pk, page]), data)


class PermissionsTests(TestCase):
    def setUp(self):
        self.controle = creer_controle()
        self.chef_user, self.chef_agent = creer_utilisateur_agent("chef", "AG-CHEF")
        self.agent_user, self.agent_agent = creer_utilisateur_agent("agent", "AG-AGENT")
        self.exterieur_user, _ = creer_utilisateur_agent("exterieur", "AG-EXT")
        MembreGroupeControle.objects.create(controle=self.controle, agent=self.chef_agent, role=RoleMission.CHEF)
        MembreGroupeControle.objects.create(controle=self.controle, agent=self.agent_agent, role=RoleMission.AGENT)

    def test_anonyme_redirige_vers_login(self):
        r = self.client.get(reverse("missions:controle_entite_detail", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r.url)

    def test_utilisateur_non_membre_refuse(self):
        self.client.login(username="exterieur", password="pass-123")
        r = self.client.get(reverse("missions:controle_entite_detail", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 403)

    def test_membre_autorise(self):
        self.client.login(username="agent", password="pass-123")
        r = self.client.get(reverse("missions:controle_entite_detail", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 200)

    def test_agent_non_chef_refuse_sur_vue_chef(self):
        self.client.login(username="agent", password="pass-123")
        r = self.client.post(reverse("missions:controle_valider", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 403)

    def test_chef_autorise_sur_vue_chef(self):
        self.client.login(username="chef", password="pass-123")
        r = self.client.post(reverse("missions:controle_valider", args=[self.controle.pk]))
        # refusé au niveau métier (statut), pas au niveau permission : donc redirect, pas 403
        self.assertEqual(r.status_code, 302)

    def test_superuser_bypass_sans_etre_membre(self):
        """Régression : request.controle doit être posé même quand la vue
        court-circuite via is_superuser (bug corrigé dans core/permissions.py)."""
        superuser = User.objects.create_superuser(username="admin", password="pass-123", email="a@a.com")
        self.client.force_login(superuser)
        r = self.client.get(reverse("missions:controle_entite_detail", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 200)

    def test_mission_create_reserve_a_l_admin(self):
        """Seule l'admin (groupe Administrateur ou superuser) peut créer une
        mission — chef et agent contrôleur, même membres d'un groupe de
        contrôle, sont refusés."""
        self.client.login(username="chef", password="pass-123")
        r = self.client.get(reverse("missions:mission_create"))
        self.assertEqual(r.status_code, 403)

    def test_mission_create_autorise_pour_superuser(self):
        superuser = User.objects.create_superuser(username="admin", password="pass-123", email="a@a.com")
        self.client.force_login(superuser)
        r = self.client.get(reverse("missions:mission_create"))
        self.assertEqual(r.status_code, 200)

    def test_mission_detail_accessible_a_un_membre_dun_controle(self):
        """La fiche mission (vue d'ensemble) est accessible à l'admin ainsi
        qu'à tout membre d'au moins un des contrôles de la mission."""
        self.client.login(username="agent", password="pass-123")
        r = self.client.get(reverse("missions:mission_detail", args=[self.controle.mission.pk]))
        self.assertEqual(r.status_code, 200)

    def test_mission_detail_refuse_a_un_tiers(self):
        self.client.login(username="exterieur", password="pass-123")
        r = self.client.get(reverse("missions:mission_detail", args=[self.controle.mission.pk]))
        self.assertEqual(r.status_code, 403)


class DashboardViewsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="admin", password="pass-123", email="a@a.com")
        self.client.force_login(self.user)

    def test_mission_list(self):
        r = self.client.get(reverse("missions:mission_list"))
        self.assertEqual(r.status_code, 200)

    def test_mission_create(self):
        entite = EntiteControlee.objects.create(nom="ACME SA")
        r = self.client.post(reverse("missions:mission_create"), {
            "entites_controlees": [entite.pk],
            "date_mission": "2026-07-10",
        })
        self.assertEqual(r.status_code, 302)
        mission = MissionControle.objects.get(entites_controlees=entite)
        controle = mission.controles_entites.get(entite=entite)
        self.assertEqual(controle.reponses.count(), 10)
        self.assertEqual(controle.statut, StatutMission.BROUILLON)

    def test_mission_create_avec_plusieurs_entites(self):
        entite_a = EntiteControlee.objects.create(nom="ACME SA")
        entite_b = EntiteControlee.objects.create(nom="Omega SARL")
        r = self.client.post(reverse("missions:mission_create"), {
            "entites_controlees": [entite_a.pk, entite_b.pk],
            "date_mission": "2026-07-10",
        })
        self.assertEqual(r.status_code, 302)
        mission = MissionControle.objects.get(entites_controlees=entite_a)
        self.assertEqual(set(mission.entites_controlees.all()), {entite_a, entite_b})
        self.assertEqual(mission.controles_entites.count(), 2)

    def test_mission_create_avec_nouvelle_entite(self):
        r = self.client.post(reverse("missions:mission_create"), {
            "nom_nouvelle_entite": "Beta SARL",
            "date_mission": "2026-07-10",
        })
        self.assertEqual(r.status_code, 302)
        mission = MissionControle.objects.get(entites_controlees__nom="Beta SARL")
        controle = mission.controles_entites.get(entite__nom="Beta SARL")
        self.assertEqual(controle.reponses.count(), 10)

    def test_mission_create_refuse_sans_entite_ni_nom(self):
        r = self.client.post(reverse("missions:mission_create"), {
            "date_mission": "2026-07-10",
        })
        self.assertEqual(r.status_code, 200)
        self.assertFalse(MissionControle.objects.exists())

    def test_mission_create_get_affiche_formulaire(self):
        r = self.client.get(reverse("missions:mission_create"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("form", r.context)

    def test_entite_ajouter(self):
        controle = creer_controle()
        mission = controle.mission
        nouvelle_entite = EntiteControlee.objects.create(nom="Omega SARL")

        r = self.client.post(reverse("missions:entite_ajouter", args=[mission.pk]), {
            "entite_controlee": nouvelle_entite.pk,
        })
        self.assertEqual(r.status_code, 302)
        self.assertEqual(mission.controles_entites.count(), 2)
        self.assertTrue(mission.controles_entites.filter(entite=nouvelle_entite).exists())

    def test_entite_ajouter_avec_nouvelle_entite(self):
        controle = creer_controle()
        mission = controle.mission

        r = self.client.post(reverse("missions:entite_ajouter", args=[mission.pk]), {
            "nom_nouvelle_entite": "Gamma SARL",
        })
        self.assertEqual(r.status_code, 302)
        self.assertTrue(mission.controles_entites.filter(entite__nom="Gamma SARL").exists())

    def test_membre_ajouter_refuse_deuxieme_chef(self):
        controle = creer_controle()
        chef_a = AgentControleur.objects.create(external_id="AG-A", nom="Nzue", prenom="Alice")
        chef_b = AgentControleur.objects.create(external_id="AG-B", nom="Mbadinga", prenom="Paul")

        self.client.post(reverse("missions:membre_ajouter", args=[controle.pk]), {
            "agent": chef_a.pk, "role": RoleMission.CHEF,
        })
        r = self.client.post(reverse("missions:membre_ajouter", args=[controle.pk]), {
            "agent": chef_b.pk, "role": RoleMission.CHEF,
        })
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.membres_groupe.chefs().count(), 1)
        self.assertTrue(controle.membres_groupe.chefs().filter(agent=chef_a).exists())

    def test_membre_et_personne_ajouter(self):
        controle = creer_controle()
        agent = AgentControleur.objects.create(external_id="AG-001", nom="Obiang", prenom="Marie")
        personne = Personne.objects.create(nom="Test", prenom="Personne")

        r = self.client.post(reverse("missions:membre_ajouter", args=[controle.pk]), {
            "agent": agent.pk, "role": RoleMission.CHEF,
        })
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.membres_groupe.count(), 1)

        r = self.client.post(reverse("missions:personne_ajouter", args=[controle.pk]), {"personne": personne.pk})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.personnes_interrogees.count(), 1)

    def test_personne_ajouter_cree_une_nouvelle_personne(self):
        controle = creer_controle()
        r = self.client.post(reverse("missions:personne_ajouter", args=[controle.pk]), {
            "nom": "Ndong", "prenom": "Alice", "email": "alice@example.com", "telephone": "01020304",
            "poste_snapshot": "Comptable", "service_snapshot": "Finance",
        })
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.personnes_interrogees.count(), 1)
        interrogee = controle.personnes_interrogees.first()
        self.assertEqual(interrogee.personne.nom, "Ndong")
        self.assertEqual(interrogee.personne.prenom, "Alice")
        self.assertEqual(interrogee.poste_snapshot, "Comptable")
        self.assertEqual(interrogee.service_snapshot, "Finance")

    def test_personne_ajouter_refuse_sans_personne_ni_nom_prenom(self):
        controle = creer_controle()
        r = self.client.post(reverse("missions:personne_ajouter", args=[controle.pk]), {})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.personnes_interrogees.count(), 0)

    def test_personne_ajouter_refuse_si_controle_verrouille(self):
        controle = creer_controle()
        controle.statut = StatutMission.PV_GENERE
        controle.save()
        personne = Personne.objects.create(nom="Test", prenom="Personne")

        r = self.client.post(reverse("missions:personne_ajouter", args=[controle.pk]), {"personne": personne.pk})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.personnes_interrogees.count(), 0)

    def test_observations_modifier(self):
        controle = creer_controle()
        r = self.client.post(reverse("missions:observations_modifier", args=[controle.pk]), {
            "commentaires_observations": "RAS",
        })
        self.assertEqual(r.status_code, 302)
        controle.refresh_from_db()
        self.assertEqual(controle.commentaires_observations, "RAS")

    def test_membre_ajouter_formulaire_invalide(self):
        controle = creer_controle()
        r = self.client.post(reverse("missions:membre_ajouter", args=[controle.pk]), {"role": RoleMission.CHEF})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(controle.membres_groupe.count(), 0)

    def test_personne_ajouter_reprend_fonction_existante(self):
        from personnes.models import Fonction, Personne

        controle = creer_controle()
        personne = Personne.objects.create(nom="Ndong", prenom="Paul")
        Fonction.objects.create(
            personne=personne, entite=controle.entite, poste="DAF", service="Finance",
        )

        r = self.client.post(reverse("missions:personne_ajouter", args=[controle.pk]), {"personne": personne.pk})
        self.assertEqual(r.status_code, 302)
        interrogee = controle.personnes_interrogees.get()
        self.assertEqual(interrogee.poste_snapshot, "DAF")
        self.assertEqual(interrogee.service_snapshot, "Finance")

    def test_infos_pv_modifier_formulaire_invalide(self):
        controle = creer_controle()
        r = self.client.post(reverse("missions:infos_pv_modifier", args=[controle.pk]), {
            "date_signature": "date-invalide",
        })
        self.assertEqual(r.status_code, 302)
        controle.refresh_from_db()
        self.assertIsNone(controle.date_signature)


class QuestionnaireFlowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="admin", password="pass-123", email="a@a.com")
        self.client.force_login(self.user)
        self.controle = creer_controle()
        declarer_tous_traitements(self.controle)

    def test_parcours_des_5_pages_complete_le_statut(self):
        for page in range(1, 6):
            r = completer_page(self.client, self.controle, page)
            self.assertEqual(r.status_code, 302, r.content[:1000])

        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.QUESTIONNAIRE_COMPLETE)
        self.assertEqual(
            list(self.controle.journal.values_list("type_action", flat=True)),
            ["questionnaire_complete"],
        )

    def test_page_inconnue_renvoie_404(self):
        r = self.client.get(reverse("missions:questionnaire_page", args=[self.controle.pk, 6]))
        self.assertEqual(r.status_code, 404)

    def test_formulaire_invalide_naffiche_pas_de_redirection(self):
        r = self.client.get(reverse("missions:questionnaire_page", args=[self.controle.pk, 1]))
        formset = r.context["formset"]
        data = {f"form-{k}": v for k, v in formset.management_form.initial.items()}
        # TOTAL_FORMS cassé -> formset.is_valid() est False sans même regarder les données des lignes.
        data["form-TOTAL_FORMS"] = "999"
        r = self.client.post(reverse("missions:questionnaire_page", args=[self.controle.pk, 1]), data)
        self.assertEqual(r.status_code, 200)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.BROUILLON)


class QuestionnaireChecklistTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="admin", password="pass-123", email="a@a.com")
        self.client.force_login(self.user)
        self.controle = creer_controle()

    def _donnees_checklist(self, codes_declares):
        r = self.client.get(reverse("missions:questionnaire_checklist", args=[self.controle.pk]))
        formset = r.context["formset"]
        lignes = r.context["lignes"]
        data = {f"form-{k}": v for k, v in formset.management_form.initial.items()}
        for i, (form, page1) in enumerate(lignes):
            data[f"form-{i}-id"] = page1.pk
            if page1.reponse.traitement in codes_declares:
                data[f"form-{i}-declaration_effectuee"] = "on"
        return data

    def test_aucun_traitement_declare_par_defaut(self):
        r = self.client.get(reverse("missions:questionnaire_page", args=[self.controle.pk, 1]))
        self.assertEqual(list(r.context["lignes"]), [])

    def test_checklist_filtre_les_traitements_de_la_page_1(self):
        data = self._donnees_checklist({"a", "b"})
        r = self.client.post(reverse("missions:questionnaire_checklist", args=[self.controle.pk]), data)
        self.assertRedirects(r, reverse("missions:questionnaire_page", args=[self.controle.pk, 1]))

        r = self.client.get(reverse("missions:questionnaire_page", args=[self.controle.pk, 1]))
        codes = {ligne.reponse.traitement for _, ligne in r.context["lignes"]}
        self.assertEqual(codes, {"a", "b"})

    def test_traitements_non_declares_classes_non_conformes(self):
        data = self._donnees_checklist({"a"})
        self.client.post(reverse("missions:questionnaire_checklist", args=[self.controle.pk]), data)

        self.controle.refresh_from_db()
        for reponse in self.controle.reponses.all():
            if reponse.traitement == "a":
                self.assertEqual(reponse.evaluation, "")
            else:
                self.assertEqual(reponse.evaluation, EvaluationConformite.NC)

    def test_re_declarer_ne_touche_pas_a_une_evaluation_deja_forcee(self):
        """Une fois classé NC par la checkliste, cocher le traitement comme
        déclaré ensuite ne doit pas modifier l'évaluation déjà enregistrée —
        seule l'évaluation manuelle (page dédiée) le fait."""
        self.client.post(
            reverse("missions:questionnaire_checklist", args=[self.controle.pk]),
            self._donnees_checklist(set()),
        )
        reponse_a = self.controle.reponses.get(traitement="a")
        self.assertEqual(reponse_a.evaluation, EvaluationConformite.NC)

        self.client.post(
            reverse("missions:questionnaire_checklist", args=[self.controle.pk]),
            self._donnees_checklist({"a"}),
        )
        reponse_a.refresh_from_db()
        self.assertEqual(reponse_a.evaluation, EvaluationConformite.NC)


@override_settings(MEDIA_ROOT=MEDIA_ROOT_TEST)
class ClotureControleTests(TestCase):
    def setUp(self):
        self.chef_user, self.chef_agent = creer_utilisateur_agent("chef", "AG-CHEF")
        self.controle = creer_controle()
        MembreGroupeControle.objects.create(controle=self.controle, agent=self.chef_agent, role=RoleMission.CHEF)
        self.client.login(username="chef", password="pass-123")

    def test_marquer_pv_genere_refuse_si_questionnaire_incomplet(self):
        r = self.client.post(reverse("missions:pv_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.BROUILLON)

    def test_parcours_complet_jusqu_a_validation(self):
        for page in range(1, 6):
            completer_page(self.client, self.controle, page)

        r = self.client.post(reverse("missions:pv_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_GENERE)
        self.assertTrue(self.controle.est_verrouillee)
        self.assertTrue(self.controle.pv_document.name)
        self.assertTrue(self.controle.pv_document.name.endswith(".docx"))

        scan = SimpleUploadedFile("scan.pdf", b"%PDF-1.4 contenu factice", content_type="application/pdf")
        r = self.client.post(reverse("missions:scan_uploader", args=[self.controle.pk]), {"scan_signe": scan})
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_SCAN_UPLOAD)
        self.assertTrue(self.controle.scan_signe.name)

        r = self.client.post(reverse("missions:rapport_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.RAPPORT_GENERE)

        rapport_scan = SimpleUploadedFile("rapport.pdf", b"%PDF-1.4 contenu factice", content_type="application/pdf")
        r = self.client.post(
            reverse("missions:rapport_uploader", args=[self.controle.pk]), {"rapport_signe": rapport_scan}
        )
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.RAPPORT_SCAN_UPLOAD)
        self.assertTrue(self.controle.rapport_signe.name)

        r = self.client.post(reverse("missions:controle_valider", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.VALIDEE)

    def test_rapport_uploader_refuse_extension_non_autorisee(self):
        self.controle.statut = StatutMission.RAPPORT_GENERE
        self.controle.save()

        rapport = SimpleUploadedFile("rapport.exe", b"binaire", content_type="application/octet-stream")
        r = self.client.post(
            reverse("missions:rapport_uploader", args=[self.controle.pk]), {"rapport_signe": rapport}
        )
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.RAPPORT_GENERE)
        self.assertFalse(self.controle.rapport_signe)

    def test_valider_refuse_sans_rapport_signe_uploade(self):
        self.controle.statut = StatutMission.RAPPORT_GENERE
        self.controle.save()

        r = self.client.post(reverse("missions:controle_valider", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.RAPPORT_GENERE)

    def test_rapport_marquer_genere_refuse_avant_scan_du_pv(self):
        self.controle.statut = StatutMission.PV_GENERE
        self.controle.save()

        r = self.client.post(reverse("missions:rapport_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_GENERE)

    def test_scan_refuse_extension_non_autorisee(self):
        self.controle.statut = StatutMission.PV_GENERE
        self.controle.save()

        scan = SimpleUploadedFile("scan.exe", b"binaire", content_type="application/octet-stream")
        r = self.client.post(reverse("missions:scan_uploader", args=[self.controle.pk]), {"scan_signe": scan})
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_GENERE)
        self.assertFalse(self.controle.scan_signe)

    def test_valider_refuse_sans_rapport_genere(self):
        self.controle.statut = StatutMission.PV_SCAN_UPLOAD
        self.controle.save()

        r = self.client.post(reverse("missions:controle_valider", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_SCAN_UPLOAD)

    def test_pv_marquer_genere_refuse_si_deja_genere(self):
        self.controle.statut = StatutMission.PV_GENERE
        self.controle.save()

        r = self.client.post(reverse("missions:pv_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_GENERE)

    def test_rapport_marquer_genere_refuse_si_deja_genere(self):
        self.controle.statut = StatutMission.RAPPORT_GENERE
        self.controle.save()

        r = self.client.post(reverse("missions:rapport_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.RAPPORT_GENERE)

    def test_scan_uploader_refuse_avant_pv_genere(self):
        scan = SimpleUploadedFile("scan.pdf", b"%PDF-1.4", content_type="application/pdf")
        r = self.client.post(reverse("missions:scan_uploader", args=[self.controle.pk]), {"scan_signe": scan})
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.BROUILLON)
        self.assertFalse(self.controle.scan_signe)

    def test_rapport_uploader_refuse_avant_rapport_genere(self):
        self.controle.statut = StatutMission.PV_SCAN_UPLOAD
        self.controle.save()

        rapport = SimpleUploadedFile("rapport.pdf", b"%PDF-1.4", content_type="application/pdf")
        r = self.client.post(
            reverse("missions:rapport_uploader", args=[self.controle.pk]), {"rapport_signe": rapport}
        )
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_SCAN_UPLOAD)
        self.assertFalse(self.controle.rapport_signe)

    @patch("missions.views.generate_pv")
    def test_pv_marquer_genere_echec_generation_affiche_message_et_ne_verrouille_pas(self, mock_generate_pv):
        mock_generate_pv.side_effect = RuntimeError("boom")
        self.controle.statut = StatutMission.QUESTIONNAIRE_COMPLETE
        self.controle.save()

        r = self.client.post(reverse("missions:pv_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.QUESTIONNAIRE_COMPLETE)
        self.assertFalse(self.controle.est_verrouillee)
        self.assertFalse(self.controle.pv_document)

    @patch("missions.views.generate_pv")
    def test_pv_marquer_genere_sauvegarde_le_pdf_si_disponible(self, mock_generate_pv):
        with tempfile.TemporaryDirectory() as dossier:
            chemin_docx = os.path.join(dossier, "pv.docx")
            chemin_pdf = os.path.join(dossier, "pv.pdf")
            with open(chemin_docx, "wb") as f:
                f.write(b"docx factice")
            with open(chemin_pdf, "wb") as f:
                f.write(b"%PDF-1.4 factice")
            mock_generate_pv.return_value = {"docx": chemin_docx, "pdf": chemin_pdf}

            self.controle.statut = StatutMission.QUESTIONNAIRE_COMPLETE
            self.controle.save()
            r = self.client.post(reverse("missions:pv_marquer_genere", args=[self.controle.pk]))

        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.statut, StatutMission.PV_GENERE)
        self.assertTrue(self.controle.pv_document.name)
        self.assertTrue(self.controle.pv_document_pdf.name)


@override_settings(MEDIA_ROOT=MEDIA_ROOT_TEST)
class EvaluationEtInfosPVTests(TestCase):
    def setUp(self):
        self.chef_user, self.chef_agent = creer_utilisateur_agent("chef", "AG-CHEF")
        self.controle = creer_controle()
        MembreGroupeControle.objects.create(controle=self.controle, agent=self.chef_agent, role=RoleMission.CHEF)
        self.client.login(username="chef", password="pass-123")

    def test_agent_non_chef_refuse_sur_evaluation(self):
        agent_user, agent_agent = creer_utilisateur_agent("agent", "AG-AGENT")
        MembreGroupeControle.objects.create(controle=self.controle, agent=agent_agent, role=RoleMission.AGENT)
        self.client.login(username="agent", password="pass-123")

        r = self.client.get(reverse("missions:evaluation_page", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 403)

    def test_enregistrer_evaluation(self):
        r = self.client.get(reverse("missions:evaluation_page", args=[self.controle.pk]))
        formset = r.context["formset"]
        data = {f"form-{k}": v for k, v in formset.management_form.initial.items()}
        for i, form in enumerate(formset.forms):
            data[f"form-{i}-id"] = form.instance.pk
            data[f"form-{i}-observations_controleur"] = ""
            data[f"form-{i}-observations_entite"] = ""
        # Ne renseigne un verdict que pour la première ligne (traitement "a").
        data["form-0-evaluation"] = EvaluationConformite.CTO
        data["form-0-observations_controleur"] = "RAS."

        r = self.client.post(reverse("missions:evaluation_page", args=[self.controle.pk]), data)
        self.assertEqual(r.status_code, 302, r.content[:1000])

        reponse_a = self.controle.reponses.get(traitement="a")
        self.assertEqual(reponse_a.evaluation, EvaluationConformite.CTO)
        self.assertEqual(reponse_a.observations_controleur, "RAS.")

    def test_evaluation_formulaire_invalide(self):
        r = self.client.get(reverse("missions:evaluation_page", args=[self.controle.pk]))
        formset = r.context["formset"]
        data = {f"form-{k}": v for k, v in formset.management_form.initial.items()}
        for i, form in enumerate(formset.forms):
            data[f"form-{i}-id"] = form.instance.pk
            data[f"form-{i}-observations_controleur"] = ""
            data[f"form-{i}-observations_entite"] = ""
        data["form-0-evaluation"] = "verdict-invalide"

        r = self.client.post(reverse("missions:evaluation_page", args=[self.controle.pk]), data)
        self.assertEqual(r.status_code, 200)
        reponse_a = self.controle.reponses.get(traitement="a")
        self.assertEqual(reponse_a.evaluation, "")

    def test_evaluation_impossible_si_controle_verrouille(self):
        self.controle.statut = StatutMission.PV_GENERE
        self.controle.save()

        reponse = self.controle.reponses.first()
        reponse.evaluation = EvaluationConformite.CTO
        with self.assertRaises(ValidationError):
            reponse.save()

    def test_infos_pv_modifier(self):
        r = self.client.post(reverse("missions:infos_pv_modifier", args=[self.controle.pk]), {
            "mode_pv": "in situ",
            "nom_representant_entite": "M. Test Représentant",
            "heure_controle": "09h30",
            "deliberation_numero": "001/2026",
            "deliberation_organe": "Conseil de l'APDPVP",
            "lieu_signature": "Libreville",
            "date_signature": "2026-07-10",
            "heure_signature": "12h00",
        })
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()
        self.assertEqual(self.controle.nom_representant_entite, "M. Test Représentant")
        self.assertEqual(self.controle.mode_pv, "in situ")


@override_settings(MEDIA_ROOT=MEDIA_ROOT_TEST)
class GenerationPVTests(TestCase):
    """Vérifie l'intégration réelle avec missions/generate_pv.py (pas de mock)."""

    def setUp(self):
        self.chef_user, self.chef_agent = creer_utilisateur_agent("chef", "AG-CHEF")
        self.controle = creer_controle()
        MembreGroupeControle.objects.create(controle=self.controle, agent=self.chef_agent, role=RoleMission.CHEF)
        self.controle.nom_representant_entite = "M. Jean Test"
        self.controle.mode_pv = "in situ"
        self.controle.save()
        self.controle.statut = StatutMission.QUESTIONNAIRE_COMPLETE
        self.controle.save()
        self.client.login(username="chef", password="pass-123")

    def test_pv_document_contient_les_donnees_du_controle(self):
        from docx import Document

        reponse = self.controle.reponses.get(traitement="a")
        reponse.evaluation = EvaluationConformite.CTO
        reponse.observations_controleur = "Observation spécifique au traitement a."
        reponse.save()

        r = self.client.post(reverse("missions:pv_marquer_genere", args=[self.controle.pk]))
        self.assertEqual(r.status_code, 302)
        self.controle.refresh_from_db()

        self.assertTrue(self.controle.pv_document.name)
        with self.controle.pv_document.open("rb") as f:
            doc = Document(f)
        texte_complet = "\n".join(p.text for p in doc.paragraphs)
        texte_complet += "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

        self.assertIn("ACME SA", texte_complet)
        self.assertIn("Jean Test", texte_complet)
        self.assertIn("Observation spécifique au traitement a.", texte_complet)
